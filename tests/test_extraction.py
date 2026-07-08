"""Tests for the extraction pipeline (no API key or network needed)."""

import json
from types import SimpleNamespace

import pytest

from backend import database, exports, template_store
from backend.extractor import extract_data, extract_pdf_text
from backend.schema import THINFILM_SCHEMA

SAMPLE_RESULT = {
    "paper": {
        "title": "ZnO thin films by RF sputtering",
        "authors": ["A. Author", "B. Author"],
        "journal": "Thin Solid Films",
        "year": 2024,
        "doi": "10.1000/example",
    },
    "samples": [
        {
            "material": "ZnO",
            "substrate": "glass",
            "deposition_method": "RF magnetron sputtering",
            "deposition_temperature_c": 200,
            "annealing_temperature_c": 400,
            "thickness_nm": 150,
            "bandgap_ev": 3.28,
            "resistivity_ohm_cm": 0.004,
            "transmittance_percent": 85,
            "other_properties": [
                {"name": "carrier concentration", "value": "2.1e20", "unit": "cm^-3"}
            ],
        }
    ],
    "notes": None,
}

# The extractor/retry tests exercise parse + validation logic, not the shipped
# thin-film schema. Pin them to a small legacy template that SAMPLE_RESULT
# satisfies, so they're independent of changes to thin_film_deposition.json.
LEGACY_TF_TEMPLATE = {
    "name": "legacy_test",
    "record_label": "sample",
    "paper_fields": ["title", "authors", "journal", "year", "doi"],
    "include_other_properties": True,
    "fields": [
        {"name": "material", "type": "string", "nullable": False},
        {"name": "substrate", "type": "string"},
        {"name": "deposition_method", "type": "string"},
        {"name": "deposition_temperature_c", "type": "number"},
        {"name": "annealing_temperature_c", "type": "number"},
        {"name": "thickness_nm", "type": "number"},
        {"name": "bandgap_ev", "type": "number"},
        {"name": "resistivity_ohm_cm", "type": "number"},
        {"name": "transmittance_percent", "type": "number"},
    ],
}

# A record shaped like the new 32-field literature-tracking thin-film template.
NEW_TF_RESULT = {
    "paper": {
        "title": "WO3 gasochromic H2 sensor by RF sputtering",
        "authors": ["Y. Lee", "S. Kalanur"],
        "journal": "Sens. Actuators B",
        "year": 2017,
        "doi": "10.1016/example",
    },
    "samples": [
        {
            "ref_number": "NR",
            "first_author": "Lee et al.",
            "year": "2017",
            "journal": "Sens. Actuators B",
            "paper_title": "WO3 gasochromic H2 sensor by RF sputtering",
            "sputtering_type": "RF magnetron sputtering",
            "power_w": "150",
            "target_material": "WO3",
            "target_purity_percent": "99.99",
            "target_substrate_distance_cm": "NR",
            "ar_flow_sccm": "30",
            "o2_flow_sccm": "NR",
            "base_pressure": "4e-6 Torr",
            "work_pressure": "10 mTorr",
            "rotation_speed_rpm": "NR",
            "deposition_time_min": "50",
            "deposition_rate_nm_min": "NR",
            "film_thickness_nm": "760",
            "substrate_type": "glass",
            "substrate_temp_c": "25",
            "annealing": "NR",
            "xrd_crystal_phase": "amorphous",
            "xrd_peak_2theta": "NR",
            "grain_size_nm": "NR",
            "surface_roughness_nm": "NR",
            "h2_concentration_tested": "1%",
            "optical_response_wl_shift": "transmittance change >50%",
            "response_time_s": "NR",
            "recovery_time_s": "NR",
            "hysteresis": "NR",
            "sensing_temp_c": "25",
            "main_finding": "Nano-columnar WO3-Pd films show strong gasochromic response to 1% H2.",
        }
    ],
    "notes": None,
}

NEW_TF_RECORD = {
    "id": 5,
    "filename": "wo3_paper.pdf",
    "created_at": "2026-07-07T00:00:00+00:00",
    "template": "thin_film_deposition",
    "data": NEW_TF_RESULT,
}


class FakeResponse:
    """Mimics a google-genai GenerateContentResponse."""

    def __init__(self, text: str):
        self._text = text

    @property
    def text(self) -> str:
        return self._text


class FakeClient:
    """Mimics genai.Client, returning one canned response per call
    from client.models.generate_content()."""

    def __init__(self, responses: list[str]):
        self._responses = responses
        self.calls = 0
        self.prompts: list[str] = []
        self.models = SimpleNamespace(generate_content=self._generate_content)

    def _generate_content(self, model, contents, config=None):
        self.prompts.append(contents)
        response = FakeResponse(self._responses[min(self.calls, len(self._responses) - 1)])
        self.calls += 1
        return response


# ---------------------------------------------------------------- schema


def test_schema_shape():
    assert THINFILM_SCHEMA["type"] == "object"
    assert THINFILM_SCHEMA["additionalProperties"] is False
    assert set(THINFILM_SCHEMA["required"]) == {"paper", "samples", "notes"}


def test_schema_objects_forbid_extra_properties():
    def walk(node):
        if isinstance(node, dict):
            if node.get("type") == "object" or "object" in (node.get("type") or []):
                assert node.get("additionalProperties") is False
            for value in node.values():
                walk(value)
        elif isinstance(node, list):
            for item in node:
                walk(item)

    walk(THINFILM_SCHEMA)


# ---------------------------------------------------------------- extractor


def test_extract_pdf_text_missing_file():
    with pytest.raises(FileNotFoundError):
        extract_pdf_text("does-not-exist.pdf")


def test_extract_data_parses_model_output():
    client = FakeClient([json.dumps(SAMPLE_RESULT)])
    result = extract_data("some paper text", client=client, template=LEGACY_TF_TEMPLATE)
    assert result == SAMPLE_RESULT
    assert result["samples"][0]["material"] == "ZnO"
    assert client.calls == 1


def test_extract_data_retries_once_on_invalid_json():
    client = FakeClient(["this is {not json", json.dumps(SAMPLE_RESULT)])
    result = extract_data("some paper text", client=client, template=LEGACY_TF_TEMPLATE)
    assert result == SAMPLE_RESULT
    assert client.calls == 2
    # The retry prompt must be stricter than the original
    assert "previous response was not valid JSON" in client.prompts[1]


def test_extract_data_retries_on_schema_mismatch():
    # Valid JSON, but missing required keys -> should retry, then succeed
    client = FakeClient([json.dumps({"paper": {}}), json.dumps(SAMPLE_RESULT)])
    result = extract_data("some paper text", client=client, template=LEGACY_TF_TEMPLATE)
    assert result == SAMPLE_RESULT
    assert client.calls == 2


def test_extract_data_retries_transient_api_errors(monkeypatch):
    from google.genai import errors as genai_errors

    from backend import extractor as extractor_module

    monkeypatch.setattr(extractor_module, "BACKOFF_SECONDS", (0,))  # no sleep in tests

    class FlakyClient(FakeClient):
        def __init__(self, responses):
            super().__init__(responses)
            self.failures = 1

        def _generate_content(self, model, contents, config=None):
            if self.failures:
                self.failures -= 1
                raise genai_errors.ServerError(503, {"error": {"message": "high demand"}})
            return super()._generate_content(model, contents, config)

    client = FlakyClient([json.dumps(SAMPLE_RESULT)])
    result = extract_data("some paper text", client=client, template=LEGACY_TF_TEMPLATE)
    assert result == SAMPLE_RESULT  # transient 503 was retried, not surfaced


def test_extract_data_fails_cleanly_after_retry():
    client = FakeClient(["still not json", "nope, again"])
    with pytest.raises(ValueError, match="did not return valid JSON"):
        extract_data("some paper text", client=client, template=LEGACY_TF_TEMPLATE)
    assert client.calls == 2  # exactly one retry, no infinite loop


# ---------------------------------------------------------------- templates


def test_thin_film_template_is_32_field_literature_table():
    """The thin-film template is now a 32-column literature-tracking table
    (string fields, NR for unreported) — one row per paper."""
    template = template_store.load_template("thin_film_deposition")
    labels = [label for _, label in template_store.field_labels(template)]
    assert len(labels) == 32
    assert labels[0] == "Ref#"
    assert labels[-1] == "Main Finding"
    assert "Power (W)" in labels and "Hysteresis (Y/N/NR)" in labels
    assert template["unreported_value"] == "NR"
    assert template["include_other_properties"] is False

    schema = template_store.build_schema(template)
    sample = schema["properties"]["samples"]["items"]
    # every field is a required string (so "NR" is always valid, null is not)
    assert len(sample["required"]) == 32
    assert all(sample["properties"][n]["type"] == "string" for n in sample["required"])
    assert "other_properties" not in sample["properties"]


def test_thin_film_prompt_uses_nr_convention():
    template = template_store.load_template("thin_film_deposition")
    prompt = template_store.build_system_prompt(template, template_store.build_schema(template))
    assert 'use the exact string "NR"' in prompt
    assert "use null" not in prompt


def test_nr_convention_is_scoped_to_templates_that_set_it():
    # chemistry template must be untouched — still uses null for missing values
    chem = template_store.load_template("chemical_reaction_screening")
    assert "unreported_value" not in chem
    prompt = template_store.build_system_prompt(chem, template_store.build_schema(chem))
    assert "use null for anything not" in prompt
    assert '"NR"' not in prompt


def test_list_templates_includes_shipped():
    names = {t["name"] for t in template_store.list_templates()}
    assert {"thin_film_deposition", "custom"} <= names


def test_custom_template_is_minimal():
    template = template_store.load_template("custom")
    assert template["fields"] == []
    assert template["paper_fields"] == ["title", "authors", "year"]
    schema = template_store.build_schema(template)
    paper = schema["properties"]["paper"]
    assert list(paper["properties"]) == ["title", "authors", "year"]


def test_load_unknown_template_raises():
    with pytest.raises(FileNotFoundError):
        template_store.load_template("does_not_exist")
    with pytest.raises(ValueError):
        template_store.load_template("../evil")


def test_create_template_roundtrip(tmp_path):
    spec = {
        "name": "battery_cycling",
        "display_name": "Battery cycling",
        "fields": [
            {"name": "cathode_material", "type": "string", "nullable": False,
             "description": "Cathode chemistry, e.g. NMC811"},
            {"name": "capacity_mah_g", "type": "number",
             "description": "Specific capacity in mAh/g"},
            {"name": "cycles", "type": "integer", "description": "Cycle count"},
        ],
    }
    created = template_store.create_template(spec, templates_dir=tmp_path)
    loaded = template_store.load_template("battery_cycling", templates_dir=tmp_path)
    assert loaded == created

    schema = template_store.build_schema(loaded)
    sample = schema["properties"]["samples"]["items"]
    assert sample["properties"]["cathode_material"] == {"type": "string"}
    assert sample["properties"]["capacity_mah_g"] == {"type": ["number", "null"]}
    assert sample["additionalProperties"] is False
    assert "other_properties" in sample["properties"]

    # names are unique — creating it again fails
    with pytest.raises(FileExistsError):
        template_store.create_template(spec, templates_dir=tmp_path)


def test_create_template_rejects_bad_specs(tmp_path):
    with pytest.raises(ValueError):  # bad slug
        template_store.create_template({"name": "Bad Name!", "fields": []}, tmp_path)
    with pytest.raises(ValueError):  # bad field type
        template_store.create_template(
            {"name": "x1", "fields": [{"name": "f", "type": "datetime"}]}, tmp_path
        )
    with pytest.raises(ValueError):  # duplicate field names
        template_store.create_template(
            {"name": "x2", "fields": [
                {"name": "f", "type": "string"}, {"name": "f", "type": "number"}
            ]}, tmp_path
        )
    with pytest.raises(ValueError):  # unknown paper field
        template_store.create_template(
            {"name": "x3", "fields": [], "paper_fields": ["title", "abstract"]}, tmp_path
        )


def test_extract_data_respects_custom_template(tmp_path):
    template = template_store.create_template(
        {
            "name": "cell_viability",
            "fields": [
                {"name": "cell_line", "type": "string", "nullable": False},
                {"name": "viability_percent", "type": "number"},
            ],
            "paper_fields": ["title", "authors", "year"],
        },
        templates_dir=tmp_path,
    )
    payload = {
        "paper": {"title": "T", "authors": ["A"], "year": 2020},
        "samples": [
            {"cell_line": "HeLa", "viability_percent": 87.5, "other_properties": []}
        ],
        "notes": None,
    }
    client = FakeClient([json.dumps(payload)])
    result = extract_data("some paper text", client=client, template=template)
    assert result == payload
    assert client.calls == 1

    # a thin-film-shaped payload must FAIL validation under this template
    bad_client = FakeClient([json.dumps(SAMPLE_RESULT), json.dumps(SAMPLE_RESULT)])
    with pytest.raises(ValueError, match="did not return valid JSON"):
        extract_data("some paper text", client=bad_client, template=template)
    assert bad_client.calls == 2  # retried once, then failed cleanly


# ---------------------------------------------------------------- exports

SAMPLE_RECORD = {
    "id": 1,
    "filename": "paper.pdf",
    "created_at": "2026-07-07T00:00:00+00:00",
    "data": SAMPLE_RESULT,
}


def test_export_excel_thin_film_reference_layout():
    import io

    from openpyxl import load_workbook

    content = exports.to_excel_bytes(NEW_TF_RECORD)
    assert content[:2] == b"PK"  # .xlsx is a zip container

    wb = load_workbook(io.BytesIO(content))
    assert wb.sheetnames[0] == "Literature Parameters"
    ws = wb["Literature Parameters"]

    # row 1 title, row 2 subtitle, row 3 header, row 4+ data
    assert "WO3 RF Sputtering Parameters" in ws["A1"].value
    assert "Literature Parameters" in ws["A2"].value

    header = [c.value for c in ws[3]]
    assert header[0] == "Ref#"
    assert header[-1] == "Main Finding"
    for label in ("Power (W)", "Target Material", "Film Thickness (nm)", "Sputtering Type"):
        assert label in header

    row = dict(zip(header, [c.value for c in ws[4]]))
    assert row["Ref#"] == 1  # sequential, not extracted
    assert row["Power (W)"] == "150"
    assert row["Target Material"] == "WO3"
    assert row["Sputtering Type"] == "RF magnetron sputtering"
    assert row["Film Thickness (nm)"] == "760"
    assert row["O2 Flow (sccm)"] == "NR"  # unreported -> NR, not blank


def test_export_all_excel_thin_film_one_row_per_paper():
    import io

    from openpyxl import load_workbook

    second = {**NEW_TF_RECORD, "id": 9, "filename": "other.pdf"}
    content = exports.to_excel_all_bytes([NEW_TF_RECORD, second])
    ws = load_workbook(io.BytesIO(content))["Literature Parameters"]
    # data starts at row 4; Ref# column is sequential 1..N
    refs = [ws.cell(row=r, column=1).value for r in (4, 5)]
    assert refs == [1, 2]


CHEM_RESULT = {
    "paper": {
        "title": "Suzuki coupling solvent screening",
        "authors": ["J. Sherwood"],
        "journal": "Beilstein J. Org. Chem.",
        "year": 2020,
        "doi": "10.3762/bjoc.16.89",
    },
    "samples": [
        {
            "product": "felbinac",
            "reagents": "4-bromophenylacetic acid, phenylboronic acid",
            "catalyst": "palladium acetate",
            "solvent": "2-propanol, water",
            "temperature_c": 65,
            "reaction_time_h": 20,
            "yield_percent": 93,
            "other_properties": [],
        }
    ],
    "notes": None,
}

CHEM_RECORD = {
    "id": 2,
    "filename": "chem.pdf",
    "created_at": "2026-07-07T00:00:00+00:00",
    "template": "chemical_reaction_screening",
    "data": CHEM_RESULT,
}


def test_export_excel_uses_chemistry_columns_for_chemistry_record():
    import io

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(exports.to_excel_bytes(CHEM_RECORD)))
    ws = wb["Samples"]
    header = [c.value for c in ws[1]]

    for label in ("Product", "Reagents", "Catalyst", "Solvent", "Yield (%)"):
        assert label in header
    assert "Thickness (nm)" not in header  # no thin-film columns leak in

    values = dict(zip(header, [c.value for c in ws[2]]))
    assert values["Product"] == "felbinac"
    assert values["Catalyst"] == "palladium acetate"
    assert values["Yield (%)"] == 93


def test_export_word_and_pdf_use_chemistry_fields():
    import io

    import pdfplumber
    from docx import Document

    doc = Document(io.BytesIO(exports.to_word_bytes(CHEM_RECORD)))
    cells = [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    assert "Catalyst" in cells and "palladium acetate" in cells
    assert "Thickness (nm)" not in cells
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert any(h.startswith("Reaction 1") for h in headings)  # record_label used

    with pdfplumber.open(io.BytesIO(exports.to_pdf_bytes(CHEM_RECORD))) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert "Yield (%)" in text and "felbinac" in text
    assert "Thickness (nm)" not in text


def test_export_all_excel_splits_mixed_templates_into_sheets():
    import io

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(exports.to_excel_all_bytes([NEW_TF_RECORD, CHEM_RECORD])))
    # thin-film uses the reference sheet; chemistry uses the generic flat sheet
    assert "Literature Parameters" in wb.sheetnames
    assert "chemical_reaction_screening" in wb.sheetnames

    thin_header = [c.value for c in wb["Literature Parameters"][3]]  # header is row 3
    chem_header = [c.value for c in wb["chemical_reaction_screening"][1]]
    assert "Power (W)" in thin_header and "Product" not in thin_header
    assert "Product" in chem_header and "Power (W)" not in chem_header


def test_export_word_contains_thin_film_fields():
    import io

    from docx import Document

    content = exports.to_word_bytes(NEW_TF_RECORD)
    assert content[:2] == b"PK"  # .docx is a zip container

    doc = Document(io.BytesIO(content))
    cells = [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    assert "Power (W)" in cells and "150" in cells
    assert "Sputtering Type" in cells and "RF magnetron sputtering" in cells
    assert "NR" in cells  # unreported fields carry through as NR


def test_export_pdf_contains_thin_film_fields():
    import io

    import pdfplumber

    content = exports.to_pdf_bytes(NEW_TF_RECORD)
    assert content[:5] == b"%PDF-"

    with pdfplumber.open(io.BytesIO(content)) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert "RF magnetron sputtering" in text
    assert "Power (W)" in text
    assert "Main Finding" in text


# ---------------------------------------------------------------- database

SESSION_A = "session-aaaa"
SESSION_B = "session-bbbb"


def _engine(tmp_path):
    eng = database.get_engine(f"sqlite:///{tmp_path / 'test.db'}")
    database.init_db(eng)
    return eng


def test_database_round_trip(tmp_path):
    eng = _engine(tmp_path)

    extraction_id = database.save_extraction(SESSION_A, "paper.pdf", SAMPLE_RESULT, engine=eng)
    assert isinstance(extraction_id, int)

    record = database.get_extraction(SESSION_A, extraction_id, engine=eng)
    assert record["filename"] == "paper.pdf"
    assert record["data"] == SAMPLE_RESULT

    listing = database.list_extractions(SESSION_A, engine=eng)
    assert len(listing) == 1
    assert listing[0]["id"] == extraction_id


def test_database_get_missing_returns_none(tmp_path):
    eng = _engine(tmp_path)
    assert database.get_extraction(SESSION_A, 999, engine=eng) is None


def test_database_sessions_are_isolated(tmp_path):
    """The core multi-user guarantee: one session never sees another's rows."""
    eng = _engine(tmp_path)

    a_id = database.save_extraction(SESSION_A, "a.pdf", SAMPLE_RESULT, engine=eng)
    b_id = database.save_extraction(SESSION_B, "b.pdf", SAMPLE_RESULT, engine=eng)

    # each session lists only its own
    assert [r["id"] for r in database.list_extractions(SESSION_A, engine=eng)] == [a_id]
    assert [r["id"] for r in database.list_extractions(SESSION_B, engine=eng)] == [b_id]

    # a session cannot fetch another session's record by id
    assert database.get_extraction(SESSION_A, b_id, engine=eng) is None
    assert database.get_extraction(SESSION_B, a_id, engine=eng) is None

    # nor delete it
    assert database.delete_extraction(SESSION_A, b_id, engine=eng) is False
    assert database.get_extraction(SESSION_B, b_id, engine=eng) is not None


def test_database_stores_template(tmp_path):
    eng = _engine(tmp_path)

    default_id = database.save_extraction(SESSION_A, "a.pdf", SAMPLE_RESULT, engine=eng)
    custom_id = database.save_extraction(
        SESSION_A, "b.pdf", SAMPLE_RESULT, template="custom", engine=eng
    )

    assert database.get_extraction(SESSION_A, default_id, engine=eng)["template"] == "thin_film_deposition"
    assert database.get_extraction(SESSION_A, custom_id, engine=eng)["template"] == "custom"
    listing = {r["id"]: r["template"] for r in database.list_extractions(SESSION_A, engine=eng)}
    assert listing == {default_id: "thin_film_deposition", custom_id: "custom"}


def test_database_migration_adds_columns(tmp_path):
    db = tmp_path / "old.db"
    # simulate a pre-multiuser database (no template or session_id column)
    import sqlite3

    conn = sqlite3.connect(db)
    conn.execute(
        "CREATE TABLE extractions (id INTEGER PRIMARY KEY AUTOINCREMENT, "
        "filename TEXT NOT NULL, data TEXT NOT NULL, created_at TEXT NOT NULL)"
    )
    conn.execute(
        "INSERT INTO extractions (filename, data, created_at) VALUES (?, ?, ?)",
        ("legacy.pdf", json.dumps(SAMPLE_RESULT), "2026-01-01T00:00:00+00:00"),
    )
    conn.commit()
    conn.close()

    eng = database.get_engine(f"sqlite:///{db}")
    database.init_db(eng)  # must add template + session_id columns, not crash

    # legacy rows are tagged session_id='legacy' and template=thin_film_deposition
    record = database.get_extraction("legacy", 1, engine=eng)
    assert record["filename"] == "legacy.pdf"
    assert record["template"] == "thin_film_deposition"


def test_database_delete(tmp_path):
    eng = _engine(tmp_path)
    extraction_id = database.save_extraction(SESSION_A, "paper.pdf", SAMPLE_RESULT, engine=eng)

    assert database.delete_extraction(SESSION_A, extraction_id, engine=eng) is True
    assert database.get_extraction(SESSION_A, extraction_id, engine=eng) is None
    assert database.list_extractions(SESSION_A, engine=eng) == []
    # deleting a missing row reports False rather than raising
    assert database.delete_extraction(SESSION_A, extraction_id, engine=eng) is False


# ------------------------------------------------------- /extract API (e2e)


def test_extract_api_isolation_and_daily_limit(tmp_path, monkeypatch):
    """End-to-end via the real FastAPI app (Gemini call mocked): two sessions
    are isolated, and the per-session daily limit returns the right message
    without affecting the other session."""
    import importlib

    from fastapi.testclient import TestClient

    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 'api.db'}")
    monkeypatch.setenv("RATE_LIMIT_PER_DAY", "2")
    database._engines.clear()  # force the app to build the temp-DB engine

    import backend.main as main_mod

    importlib.reload(main_mod)
    monkeypatch.setattr(main_mod, "extract_from_pdf", lambda *a, **k: SAMPLE_RESULT)

    files = {"file": ("p.pdf", b"%PDF-1.4 fake", "application/pdf")}
    with TestClient(main_mod.app) as client:
        # session A: 2 succeed, 3rd hits the daily limit
        for _ in range(2):
            assert client.post("/extract", files=files, headers={"X-Session-Id": "A"}).status_code == 200
        blocked = client.post("/extract", files=files, headers={"X-Session-Id": "A"})
        assert blocked.status_code == 429
        assert blocked.json()["detail"] == "Daily limit reached — resets at midnight UTC."

        # session B is unaffected — its own allowance is intact
        assert client.post("/extract", files=files, headers={"X-Session-Id": "B"}).status_code == 200
        assert client.get("/usage", headers={"X-Session-Id": "B"}).json()["remaining"] == 1

        # isolation: each session lists only its own; B can't fetch A's by id
        a = client.get("/extractions", headers={"X-Session-Id": "A"}).json()
        b = client.get("/extractions", headers={"X-Session-Id": "B"}).json()
        assert len(a) == 2 and len(b) == 1
        assert client.get(f"/extractions/{a[0]['id']}", headers={"X-Session-Id": "B"}).status_code == 404

    database._engines.clear()
