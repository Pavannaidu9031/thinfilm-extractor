"""Export extraction results as Excel, Word, and PDF files (returned as bytes).

Columns are derived from each record's stored template (backend/templates/),
so exports work for any research area — thin films, chemistry, and anything
created through the template builder.
"""

import io
from xml.sax.saxutils import escape

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from backend.template_store import DEFAULT_TEMPLATE, field_labels, load_template

# Paper-identity columns shared by every template (left side of flat sheets)
PAPER_COLUMNS = ["Extraction #", "Source file", "Title", "Journal", "Year", "DOI", "Record #"]

# The thin-film template renders as a literature-tracking sheet matching the
# user's reference Excel layout: title row, subtitle row, header row, one row
# per paper. The Ref# column is a sequential row index, not extracted data.
REFERENCE_TEMPLATE = "thin_film_deposition"
REFERENCE_SHEET_NAME = "Literature Parameters"
REFERENCE_TITLE = "WO3 RF Sputtering Parameters — Extracted from Uploaded Papers"
REFERENCE_REF_FIELD = "ref_number"


def _build_reference_sheet(ws, records: list[dict], template: dict) -> None:
    """Reference literature-table layout for the thin-film template."""
    fields = field_labels(template)  # (name, label), reference column order
    ncols = len(fields)
    last_col = get_column_letter(ncols)

    # row 1: title (merged), row 2: subtitle (merged)
    ws.append([REFERENCE_TITLE])
    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")

    subtitle = (
        f"Literature Parameters · {len(records)} paper"
        f"{'' if len(records) == 1 else 's'} · one row per paper · "
        f'unreported values shown as "{template.get("unreported_value", "NR")}"'
    )
    ws.append([subtitle])
    ws.merge_cells(f"A2:{last_col}2")
    ws["A2"].font = Font(italic=True, size=10, color="666666")

    # row 3: header
    ws.append([label for _, label in fields])
    for cell in ws[3]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    # row 4+: one data row per paper; Ref# is a sequential index
    for i, record in enumerate(records, 1):
        samples = record["data"].get("samples", [])
        sample = samples[0] if samples else {}
        row = []
        for name, _ in fields:
            if name == REFERENCE_REF_FIELD:
                row.append(i)
            else:
                value = sample.get(name)
                row.append(value if value not in (None, "") else "NR")
        ws.append(row)

    ws.freeze_panes = "A4"
    for idx, (name, label) in enumerate(fields, 1):
        # narrow for Ref#, wide for title/finding, medium otherwise
        if name == REFERENCE_REF_FIELD:
            width = 6
        elif name in ("paper_title", "main_finding"):
            width = 48
        else:
            width = max(14, len(label) + 2)
        ws.column_dimensions[get_column_letter(idx)].width = width


def _dash(value):
    return "—" if value is None or value == "" else value


def _fields_for_record(record: dict) -> tuple[list[tuple[str, str]], dict | None]:
    """(field_name, label) pairs for a record, from its stored template.

    Falls back to deriving fields from the record's own data keys if the
    template file has since been deleted."""
    name = record.get("template") or DEFAULT_TEMPLATE
    try:
        template = load_template(name)
        return field_labels(template), template
    except (FileNotFoundError, ValueError):
        keys: list[str] = []
        for sample in record["data"].get("samples", []):
            for key in sample:
                if key != "other_properties" and key not in keys:
                    keys.append(key)
        return [(k, k.replace("_", " ").capitalize()) for k in keys], None


def _record_label(template: dict | None) -> str:
    return (template or {}).get("record_label", "record")


def _props_summary(sample: dict) -> str:
    parts = []
    for prop in sample.get("other_properties") or []:
        unit = f" {prop['unit']}" if prop.get("unit") else ""
        parts.append(f"{prop.get('name')} = {prop.get('value')}{unit}")
    return "; ".join(parts)


def _flat_rows(record: dict, fields: list[tuple[str, str]]) -> list[list]:
    """Flat comparison rows for one extraction (one row per sample/record).

    Missing values stay as None so Excel cells are empty and numeric columns
    remain sortable/filterable."""
    data = record["data"]
    paper = data.get("paper", {})
    base = [
        record.get("id"),
        record.get("filename"),
        paper.get("title"),
        paper.get("journal"),
        paper.get("year"),
        paper.get("doi"),
    ]
    rows = []
    samples = data.get("samples", [])
    if not samples:
        rows.append(base + [None] + [None] * len(fields) + [None, data.get("notes")])
    for i, sample in enumerate(samples, 1):
        rows.append(
            base
            + [i]
            + [sample.get(key) for key, _ in fields]
            + [_props_summary(sample) or None, data.get("notes")]
        )
    return rows


def _build_workbook(records: list[dict]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    # group records by template, preserving input order
    groups: dict[str, list[dict]] = {}
    for record in records:
        groups.setdefault(record.get("template") or DEFAULT_TEMPLATE, []).append(record)

    # one sheet per template. The thin-film template uses the bespoke
    # literature-table reference layout; every other template uses the generic
    # flat comparison layout ("Samples" when it's the only group).
    for template_name, group in groups.items():
        fields, template = _fields_for_record(group[0])

        if template_name == REFERENCE_TEMPLATE and template is not None:
            ws = wb.create_sheet(REFERENCE_SHEET_NAME[:31])
            _build_reference_sheet(ws, group, template)
            continue

        title = "Samples" if len(groups) == 1 else template_name[:31]
        ws = wb.create_sheet(title)
        header = PAPER_COLUMNS + [label for _, label in fields] + ["Other properties", "Paper notes"]
        ws.append(header)
        for record in group:
            for row in _flat_rows(record, fields):
                ws.append(row)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        widths = [11, 12, 40, 22, 7, 24, 9] + [max(16, len(l) + 2) for _, l in fields] + [60, 60]
        for idx, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(idx)].width = width

    # every other_property as its own row (nothing truncated) — only when some
    # record actually has properties (thin-film literature template has none)
    has_props = any(
        (sample.get("other_properties") or [])
        for record in records
        for sample in record["data"].get("samples", [])
    )
    if has_props:
        ws_other = wb.create_sheet("Other properties")
        ws_other.append(["Extraction #", "Source file", "Template", "Record #",
                         "Property", "Value", "Unit"])
        for record in records:
            for i, sample in enumerate(record["data"].get("samples", []), 1):
                for prop in sample.get("other_properties") or []:
                    ws_other.append([
                        record.get("id"), record.get("filename"),
                        record.get("template") or DEFAULT_TEMPLATE, i,
                        prop.get("name"), prop.get("value"), prop.get("unit"),
                    ])
        for cell in ws_other[1]:
            cell.font = Font(bold=True)
        ws_other.freeze_panes = "A2"
        for idx, width in enumerate((11, 12, 24, 9, 45, 45, 14), 1):
            ws_other.column_dimensions[get_column_letter(idx)].width = width

    # paper-level metadata, one row per paper
    ws_paper = wb.create_sheet("Papers")
    ws_paper.append(["Extraction #", "Source file", "Template", "Extracted at",
                     "Title", "Authors", "Journal", "Year", "DOI", "Notes"])
    for record in records:
        paper = record["data"].get("paper", {})
        ws_paper.append([
            record.get("id"), record.get("filename"),
            record.get("template") or DEFAULT_TEMPLATE, record.get("created_at"),
            paper.get("title"), "; ".join(paper.get("authors") or []),
            paper.get("journal"), paper.get("year"), paper.get("doi"),
            record["data"].get("notes"),
        ])
    for cell in ws_paper[1]:
        cell.font = Font(bold=True)
    ws_paper.freeze_panes = "A2"
    for idx, width in enumerate((11, 12, 24, 26, 50, 50, 24, 7, 26, 80), 1):
        ws_paper.column_dimensions[get_column_letter(idx)].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_excel_bytes(record: dict) -> bytes:
    """One extraction as a flat, comparison-ready workbook."""
    return _build_workbook([record])


def to_excel_all_bytes(records: list[dict]) -> bytes:
    """All extractions in one workbook, one comparison sheet per template."""
    return _build_workbook(records)


def to_word_bytes(record: dict) -> bytes:
    """Word document: paper metadata, one section per record, notes."""
    data = record["data"]
    paper = data.get("paper", {})
    samples = data.get("samples", [])
    fields, template = _fields_for_record(record)
    label = _record_label(template).capitalize()
    primary = fields[0][0] if fields else None

    doc = Document()
    doc.add_heading(paper.get("title") or record.get("filename", "Extraction"), level=1)

    meta_bits = [
        "; ".join(paper.get("authors") or []),
        paper.get("journal"),
        str(paper.get("year")) if paper.get("year") else None,
        paper.get("doi"),
    ]
    meta_line = " · ".join(bit for bit in meta_bits if bit)
    if meta_line:
        doc.add_paragraph(meta_line)
    doc.add_paragraph(
        f"Source file: {record.get('filename')} — extracted {record.get('created_at')} "
        f"— template: {record.get('template') or DEFAULT_TEMPLATE}"
    )

    if not samples:
        doc.add_paragraph("No records were extracted from this paper.")

    for i, sample in enumerate(samples, 1):
        heading = f"{label} {i}"
        if primary and sample.get(primary):
            heading += f": {sample[primary]}"
        doc.add_heading(heading, level=2)

        table = doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for key, field_label in fields:
            row = table.add_row().cells
            row[0].text = field_label
            row[1].text = str(_dash(sample.get(key)))

        props = sample.get("other_properties") or []
        if props:
            doc.add_paragraph("Other properties:", style="Intense Quote")
            prop_table = doc.add_table(rows=1, cols=3)
            prop_table.style = "Table Grid"
            header = prop_table.rows[0].cells
            header[0].text, header[1].text, header[2].text = "Property", "Value", "Unit"
            for prop in props:
                row = prop_table.add_row().cells
                row[0].text = str(prop.get("name", ""))
                row[1].text = str(prop.get("value", ""))
                row[2].text = str(_dash(prop.get("unit")))

    if data.get("notes"):
        doc.add_heading("Notes", level=2)
        doc.add_paragraph(data["notes"])

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_pdf_bytes(record: dict) -> bytes:
    """PDF report: paper metadata, one section per record, notes."""
    data = record["data"]
    paper = data.get("paper", {})
    samples = data.get("samples", [])
    fields, template = _fields_for_record(record)
    label = _record_label(template).capitalize()
    primary = fields[0][0] if fields else None
    styles = getSampleStyleSheet()

    def para(text, style="BodyText"):
        return Paragraph(escape(str(text)), styles[style])

    story = [para(paper.get("title") or record.get("filename", "Extraction"), "Title")]

    meta_bits = [
        "; ".join(paper.get("authors") or []),
        paper.get("journal"),
        str(paper.get("year")) if paper.get("year") else None,
        paper.get("doi"),
    ]
    meta_line = " · ".join(bit for bit in meta_bits if bit)
    if meta_line:
        story.append(para(meta_line))
    story.append(
        para(
            f"Source file: {record.get('filename')} — extracted {record.get('created_at')} "
            f"— template: {record.get('template') or DEFAULT_TEMPLATE}"
        )
    )
    story.append(Spacer(1, 6 * mm))

    if not samples:
        story.append(para("No records were extracted from this paper."))

    grid = TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BACKGROUND", (0, 0), (0, -1), colors.whitesmoke),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
    ])
    prop_grid = TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
    ])

    for i, sample in enumerate(samples, 1):
        heading = f"{label} {i}"
        if primary and sample.get(primary):
            heading += f": {sample[primary]}"
        story.append(para(heading, "Heading2"))

        rows = [[para(field_label), para(_dash(sample.get(key)))] for key, field_label in fields]
        if rows:
            story.append(Table(rows, colWidths=[55 * mm, 105 * mm], style=grid))

        props = sample.get("other_properties") or []
        if props:
            story.append(Spacer(1, 3 * mm))
            prop_rows = [[para("Property"), para("Value"), para("Unit")]]
            for prop in props:
                prop_rows.append([
                    para(prop.get("name", "")),
                    para(prop.get("value", "")),
                    para(_dash(prop.get("unit"))),
                ])
            story.append(Table(prop_rows, colWidths=[60 * mm, 75 * mm, 25 * mm], style=prop_grid))
        story.append(Spacer(1, 6 * mm))

    if data.get("notes"):
        story.append(para("Notes", "Heading2"))
        story.append(para(data["notes"]))

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, title="Rock AI extraction").build(story)
    return buf.getvalue()
